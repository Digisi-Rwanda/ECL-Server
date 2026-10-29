"""Golden tests: the engine against "ECL Calculation Specification" v1.0.

`_oracle` is a second, deliberately naive implementation written line by line
from the spec's pseudo-code (python dates, plain loops). The engine must match
it to 1e-9 on every row, and a few values are also checked by hand.
"""

from __future__ import annotations

import calendar
from datetime import date

import pandas as pd
import pytest

from app.engine.ecl_pipeline import ECL_BEFORE_CAP, ECL_CAPPED, TOTAL_ECL, run_ecl
from app.engine.ecl_rules import prepare_inputs

REP = date(2025, 1, 31)
CONFIG = [
    {"name": "Property", "haircut": 0.2, "time_to_realize": 2},
    {"name": "Cash", "haircut": 0.0, "time_to_realize": 0},
]
CURVE_MONTHS = 40


def _curve() -> pd.DataFrame:
    rows = []
    for seg, base, cure in (("Retail", 0.01, 0.10), ("Hi", 0.50, 0.0)):
        for transition, bump in (("Stage 1", 0.0), ("Stage 2", 0.02), ("Stage 3", 0.0)):
            for m in range(1, CURVE_MONTHS + 1):
                marginal = (
                    min(1.0, base + bump + m / 10_000)
                    if transition != "Stage 3"
                    else (1.0 if m == 1 else 0.0)
                )
                rows.append(
                    {
                        "SEGMENT": seg,
                        "Month": m,
                        "Transition": transition,
                        "Marginal_PD": marginal,
                        "Cure_Rate": cure,
                    }
                )
    return pd.DataFrame(rows)


def _exposure() -> pd.DataFrame:
    def loan(lid, cust, seg, stage, fp, mat, amount, eir, rep=REP):
        return {
            "Loan ID": lid,
            "Customer ID": cust,
            "SEGMENT": seg,
            "Reporting Date": rep,
            "First Payment Date": fp,
            "Maturity Date": mat,
            "Staging": stage,
            "Outstanding Amount": amount,
            "Effective Interest Rate (EIR)": eir,
        }

    return pd.DataFrame(
        [
            loan(
                "L1", "C1", "Retail", "Stage 1", date(2023, 2, 15), date(2027, 6, 10), 100_000, 0.12
            ),
            loan(
                "L2", "C1", "Retail", "Stage 2", date(2024, 1, 31), date(2025, 8, 31), 60_000, 0.18
            ),
            loan("L3", "C2", "Retail", "Stage 3", date(2022, 5, 1), date(2026, 5, 1), 30_000, 0.15),
            loan(
                "L4", "C3", "Retail", "Stage 1", date(2023, 1, 1), date(2025, 1, 5), 9_000, 0.10
            ),  # matured (same month)
            loan(
                "L5", "C3", "Retail", "Stage 2", date(2024, 6, 30), date(2026, 6, 30), 12_000, 0.0
            ),  # zero rate
            loan(
                "L6", "C4", "Retail", "Stage 2", date(2021, 9, 1), date(2027, 12, 31), 50_000, 0.09
            ),  # runs past the curve
            loan(
                "L7", "C5", "Hi", "Stage 2", date(2024, 1, 1), date(2026, 1, 31), 10_000, 0.20
            ),  # capped
            loan("L8", "C6", "Retail", "Stage 1", date(2024, 1, 1), date(2027, 1, 1), 5_000, 0.10),
            loan(
                "L8", "C6", "Retail", "Stage 1", date(2024, 1, 1), date(2027, 1, 1), 5_000, 0.10
            ),  # duplicate
            loan(
                "L9", "C7", "Retail", "Stage 1", date(2024, 1, 1), date(2027, 1, 1), 7_000, 12.0
            ),  # EIR as percent
            loan(
                "L10",
                "C8",
                "Retail",
                "Stage 1",
                date(2024, 1, 1),
                date(2027, 1, 1),
                8_000,
                0.1,
                rep=date(2024, 12, 31),
            ),
            loan(
                "L11", "C9", "Retail", "Stage 4", date(2024, 1, 1), date(2027, 1, 1), 4_000, 0.1
            ),  # bad stage
            loan(
                "L12", "C9", "Retail", "Stage 1", date(2025, 2, 1), date(2027, 1, 1), 4_000, 0.1
            ),  # FP after reporting
        ]
    )


def _collateral() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Customer ID": "C1",
                "Loan ID": "L1",
                "Outstanding Amount": 100_000,
                "Effective Interest Rate (EIR)": 0.12,
                "Property": 50_000,
                "Cash": None,
                "Gold": 999,
            },
            {
                "Customer ID": "C1",
                "Loan ID": "L2",
                "Outstanding Amount": 60_000,
                "Effective Interest Rate (EIR)": 0.18,
                "Property": 40_000,
                "Cash": 0,
                "Gold": 999,
            },  # inconsistent → lowest 40,000
            {
                "Customer ID": "WRONG",
                "Loan ID": "L3",
                "Outstanding Amount": 30_000,
                "Effective Interest Rate (EIR)": 0.15,
                "Property": 10_000,
                "Cash": -5,
                "Gold": 0,
            },  # wrong customer, negative cash
            {
                "Customer ID": "C4",
                "Loan ID": "L6",
                "Outstanding Amount": 50_000,
                "Effective Interest Rate (EIR)": None,
                "Property": 0,
                "Cash": 20_000,
                "Gold": 0,
            },  # missing EIR → exposure EIR
            {
                "Customer ID": "C99",
                "Loan ID": "L99",
                "Outstanding Amount": 1,
                "Effective Interest Rate (EIR)": 0.1,
                "Property": 1,
                "Cash": 1,
                "Gold": 1,
            },  # orphan
        ]
    )


# ── Oracle: literal spec pseudo-code ─────────────────────────────────────────


def _months_between(a: date, b: date) -> int:
    return (b.year - a.year) * 12 + (b.month - a.month)


def _add_months(d: date, n: int) -> date:
    total = d.year * 12 + d.month - 1 + n
    y, m = divmod(total, 12)
    return date(y, m + 1, min(d.day, calendar.monthrange(y, m + 1)[1]))


def _month_end(d: date) -> date:
    return date(d.year, d.month, calendar.monthrange(d.year, d.month)[1])


def _oracle(loan: dict, collateral: float, curve: pd.DataFrame) -> list[dict]:
    rep, fp, mat = loan["rep"], loan["fp"], loan["mat"]
    stage, eir, amount, seg = loan["stage"], loan["eir"], loan["amount"], loan["seg"]
    adj = _month_end(mat)
    matured = adj <= rep or (adj.year, adj.month) == (rep.year, rep.month)
    dates = [rep]
    if matured or eir == 0:
        max_n = 0
    elif stage == "Stage 1":
        max_n = 12
    elif stage == "Stage 2":
        max_n = 10**6
    else:
        max_n = 0
    n = 1
    while n <= max_n and _add_months(rep, n) <= adj:
        dates.append(_add_months(rep, n))
        n += 1
    r = eir / 12
    remaining = _months_between(fp, adj) - _months_between(fp, rep)
    if stage == "Stage 3" or matured or r == 0:
        pmt = 0.0
    else:
        pmt = (amount * r) / (1 - (1 + r) ** (-remaining))
    bal: list[float] = []
    for i, _ in enumerate(dates):
        bal.append(amount if i == 0 else max(0.0, bal[-1] * (1 + eir / 12) - pmt))
    cure = float(curve.loc[curve["SEGMENT"] == seg, "Cure_Rate"].iloc[0])
    out = []
    for t, snap in enumerate(dates):
        k = min(t, 3)
        ead = bal[t - k] * (1 + eir / 12) ** k
        pso = _months_between(fp, snap)
        sel = curve[(curve["SEGMENT"] == seg) & (curve["Transition"] == stage)]
        month = min(pso, int(sel["Month"].max()))
        pd_ = float(sel.loc[sel["Month"] == month, "Marginal_PD"].iloc[0])
        if t == 0 and stage in ("Stage 1", "Stage 2") and not (matured or eir == 0):
            pd_ = 0.0
        if stage == "Stage 3":
            pd_ = 1.0
        lgw = 0.0 if ead == 0 else max(0.0, 1 - collateral / ead)
        lgd = lgw * (1 - cure)
        loss = ead * pd_ * lgd
        disc = loss * (1 + eir / 12) ** (-t) if t > 0 else loss
        out.append(
            {
                "snap": snap,
                "pso": pso,
                "pmt": pmt,
                "bar": bal[t],
                "ead": ead,
                "pd": pd_,
                "lgw": lgw,
                "lgd": lgd,
                "loss": loss,
                "disc": disc,
            }
        )
    return out


@pytest.fixture(scope="module")
def result():
    prepared = prepare_inputs(_exposure(), _collateral(), CONFIG, REP)
    return run_ecl(prepared, _curve())


def test_exclusions_and_actions(result) -> None:
    prepared = result.prepared
    assert sorted(prepared.exposure["Loan ID"]) == ["L1", "L2", "L3", "L4", "L5", "L6", "L7"]
    assert sorted(set(prepared.excluded["Loan ID"])) == ["L10", "L11", "L12", "L8", "L9"]
    actions = {
        (e.dataset, e.loan_id or e.customer_id, e.field): e.action for e in result.exceptions
    }
    assert actions[("Exposure", "L8", "Loan ID")] == "Loan excluded"
    assert actions[("Exposure", "L9", "EIR")] == "Loan excluded"
    assert actions[("Exposure", "L10", "Reporting Date")] == "Loan excluded"
    assert actions[("Exposure", "L11", "Stage")] == "Loan excluded"
    assert actions[("Exposure", "L12", "First Payment Date")] == "Loan excluded"
    assert actions[("Collateral", "L99", "Loan ID")] == "Row ignored"
    assert actions[("Collateral", "L3", "Customer ID")] == "Exposure Customer ID used"
    assert actions[("Collateral", "L3", "Cash")] == "Treated as 0"
    assert actions[("Collateral", "C1", "Property")] == "Lowest value used"
    assert actions[("Collateral Parameters", "", "Collateral Type")] == "Treated as 0"  # Gold
    assert actions[("Collateral", "L6", "EIR")].startswith("Exposure EIR used")


def test_step1_discounted_collateral_by_hand(result) -> None:
    lgd = result.lgd.set_index("Loan ID")
    # C1: exposure 160,000; Property = lowest 40,000; Cash blank/0.
    assert lgd.loc["L1", "Total Customer Exposure"] == pytest.approx(160_000)
    assert lgd.loc["L1", "Loan Proportion"] == pytest.approx(100_000 / 160_000)
    assert lgd.loc["L1", "Allocated Property"] == pytest.approx(25_000)
    assert lgd.loc["L1", "Discounted Collateral"] == pytest.approx(25_000 * 0.8 * 1.12**-2)
    assert lgd.loc["L2", "Discounted Collateral"] == pytest.approx(15_000 * 0.8 * 1.18**-2)
    # L3 negative cash → 0; L6 cash 20,000 with no haircut and 0 years.
    assert lgd.loc["L3", "Discounted Collateral"] == pytest.approx(10_000 * 0.8 * 1.15**-2)
    assert lgd.loc["L6", "Discounted Collateral"] == pytest.approx(20_000)
    # No collateral row → 0; Gold (no parameters) never counts.
    assert lgd.loc["L7", "Discounted Collateral"] == 0
    assert (result.lgd["Discounted Gold"] == 0).all()


def test_rows_match_spec_oracle(result) -> None:
    exp = result.prepared.exposure.set_index("Loan ID")
    collateral = result.lgd.set_index("Loan ID")["Discounted Collateral"]
    curve = _curve()
    for loan_id, rows in result.rows.groupby("Loan ID"):
        e = exp.loc[loan_id]
        expected = _oracle(
            {
                "rep": e["Reporting Date"].date(),
                "fp": e["First Payment Date"].date(),
                "mat": e["Maturity Date"].date(),
                "stage": e["Stage"],
                "eir": e["EIR"],
                "amount": e["Outstanding Amount"],
                "seg": e["SEGMENT"],
            },
            float(collateral.loc[loan_id]),
            curve,
        )
        assert len(rows) == len(expected), loan_id
        for (_, got), want in zip(rows.iterrows(), expected, strict=True):
            assert got["Snapshot Date"].date() == want["snap"]
            assert got["Period Since Origination"] == want["pso"]
            for col, key in (
                ("Monthly Instalment", "pmt"),
                ("Balance After Repayment", "bar"),
                ("Balance After Missed Payment", "ead"),
                ("Marginal PD", "pd"),
                ("LGW", "lgw"),
                ("LGD", "lgd"),
                ("Credit Loss", "loss"),
                ("Discounted ECL", "disc"),
            ):
                assert got[col] == pytest.approx(want[key], rel=1e-9, abs=1e-9), (loan_id, col)


def test_treatments_and_pd_rules(result) -> None:
    rows = result.rows
    per = rows.groupby("Loan ID").size()
    assert per["L1"] == 13  # Stage 1: 12 months after the reporting row
    assert per["L2"] == 8  # Stage 2: Jan..Aug 2025
    assert per["L3"] == per["L4"] == per["L5"] == 1
    first = rows.groupby("Loan ID").first()
    assert first.loc["L1", "Marginal PD"] == 0  # reporting-date row carries no PD
    assert first.loc["L3", "Marginal PD"] == 1  # Stage 3 = 100%
    assert first.loc["L4", "Treatment"].startswith("Matured")
    assert first.loc["L4", "Marginal PD"] > 0  # matured keeps its joined PD
    assert first.loc["L5", "Treatment"].startswith("Zero rate")
    assert (
        rows.loc[rows["Loan ID"] == "L6", "PD Rule"].iloc[1:].str.contains("last available")
    ).any()
    # Missed payments: EAD at t>=3 is Balance After Repayment[t-3] grown 3 months.
    l1 = rows[rows["Loan ID"] == "L1"].reset_index(drop=True)
    assert l1.loc[5, "Balance After Missed Payment"] == pytest.approx(
        l1.loc[2, "Balance After Repayment"] * (1 + 0.12 / 12) ** 3
    )


def test_stage3_ecl_by_hand(result) -> None:
    lgd_l3 = 10_000 * 0.8 * 1.15**-2
    expected = 30_000 * 1.0 * (1 - lgd_l3 / 30_000) * (1 - 0.10)
    summary = result.loan_summary.set_index("Loan ID")
    assert summary.loc["L3", TOTAL_ECL] == pytest.approx(expected)


def test_cap_summaries_and_reconciliation(result) -> None:
    summary = result.loan_summary.set_index("Loan ID")
    assert summary.loc["L7", ECL_CAPPED] == "Yes"
    assert summary.loc["L7", ECL_BEFORE_CAP] > 10_000
    assert summary.loc["L7", TOTAL_ECL] == pytest.approx(10_000)
    assert (summary[TOTAL_ECL] <= summary["Outstanding Amount"]).all()
    assert any(e.category == "ECL Cap" and e.loan_id == "L7" for e in result.exceptions)

    stage_total = result.stage_summary.iloc[-1]
    assert stage_total["Stage"] == "Total"
    assert stage_total["ECL"] == pytest.approx(result.total_ecl)
    assert result.segment_summary.iloc[-1]["ECL"] == pytest.approx(result.total_ecl)
    coverage = result.stage_summary.set_index("Stage")["ECL Coverage Ratio"]
    assert coverage["Stage 3"] == pytest.approx(summary.loc["L3", TOTAL_ECL] / 30_000)

    checks = {c.name: c for c in result.reconciliation}
    assert set(checks) == {
        "Loan count",
        "Outstanding Amount",
        "Row Detail coverage",
        "Collateral allocation",
        "ECL build-up",
        "Summary totals",
        "ECL limits",
    }
    for name in (
        "Loan count",
        "Outstanding Amount",
        "Row Detail coverage",
        "ECL build-up",
        "Summary totals",
        "ECL limits",
    ):
        assert checks[name].passed, name
    assert checks["Loan count"].expected == 13
    # C3 (L4, L5) and C5/C4/C2 have no partial rows; C1's lowest value is fully split → passes.
    assert checks["Collateral allocation"].passed


def test_collateral_allocation_fails_when_a_customer_loan_has_no_row() -> None:
    exposure = _exposure().iloc[:2].copy()  # L1, L2 for C1
    collateral = _collateral().iloc[:1].copy()  # only L1 has a row
    run = run_ecl(prepare_inputs(exposure, collateral, CONFIG, REP), _curve())
    check = next(c for c in run.reconciliation if c.name == "Collateral allocation")
    assert not check.passed and "C1" in check.detail
    assert any(
        e.category == "Reconciliation" and e.field == "Collateral allocation"
        for e in run.exceptions
    )


def test_reporting_date_defaults_to_most_common() -> None:
    prepared = prepare_inputs(_exposure(), _collateral(), CONFIG, None)
    assert prepared.reporting_date == pd.Timestamp(REP)


def _meta() -> dict:
    return {
        "run_at": "2025-02-01 09:00",
        "run_by": "analyst@example.com",
        "run_id": "RUN1",
        "engine_version": "v1.3.0",
        "inputs": [
            {"dataset": "Exposure", "file": "ead.xlsx", "modified": "2025-02-01", "rows": 13}
        ],
    }


def test_workbooks_follow_step_11(result) -> None:
    from io import BytesIO

    import openpyxl

    from app.engine.ecl_workbooks import build_loan_detail_workbook, build_portfolio_workbook

    wb = openpyxl.load_workbook(BytesIO(build_loan_detail_workbook(result, _meta())))
    assert wb.sheetnames == ["Run Details", "Row Detail", "Loan Summary", "Exception Report"]
    detail_header = [c.value for c in wb["Row Detail"][1]]
    for col in (
        "Total Customer Exposure",
        "Loan Proportion",
        "Allocated Property",
        "Discounted Property",
        "Discounted Collateral",
        "Balance After Missed Payment",
        "Marginal PD",
        "Discounted ECL",
    ):
        assert col in detail_header
    assert wb["Row Detail"].max_row == len(result.rows) + 1
    details = {r[0].value: r[1].value for r in wb["Run Details"].iter_rows() if r[0].value}
    assert details["Specification"] == "ECL Calculation Specification Version 1.0"
    assert details["Loans excluded"] == 6

    p = openpyxl.load_workbook(BytesIO(build_portfolio_workbook(result, _meta())))
    assert p.sheetnames == ["Run Details", "Stage Summary", "Segment Summary"]
    assert [c.value for c in p["Stage Summary"][1]] == [
        "Stage",
        "Outstanding Amount",
        "ECL",
        "ECL Coverage Ratio",
    ]
    assert p["Stage Summary"]["A5"].value == "Total"


def test_row_detail_splits_loan_by_loan(result) -> None:
    from io import BytesIO

    import openpyxl

    from app.engine.ecl_workbooks import build_loan_detail_workbook

    wb = openpyxl.load_workbook(BytesIO(build_loan_detail_workbook(result, _meta(), max_rows=40)))
    tabs = [n for n in wb.sheetnames if n.startswith("Row Detail")]
    assert (
        tabs[:2] == ["Row Detail", "Row Detail 2"]
        and wb.sheetnames.index("Loan Summary") == len(tabs) + 1
    )
    seen: dict[str, str] = {}
    total = 0
    for name in tabs:
        ws = wb[name]
        assert ws.max_row <= 40 and ws["A1"].value == "Loan ID"
        for (loan,) in ws.iter_rows(min_row=2, max_col=1, values_only=True):
            assert seen.setdefault(loan, name) == name  # never split across tabs
            total += 1
    assert total == len(result.rows)


def test_exception_report_says_no_exceptions_when_clean() -> None:
    from io import BytesIO

    import openpyxl

    from app.engine.ecl_workbooks import build_loan_detail_workbook

    exposure = _exposure().iloc[[0]]
    collateral = _collateral().iloc[[0]].drop(columns=["Gold", "Cash"])
    run = run_ecl(prepare_inputs(exposure, collateral, CONFIG, REP), _curve())
    assert run.exceptions == []
    wb = openpyxl.load_workbook(BytesIO(build_loan_detail_workbook(run, _meta())))
    ws = wb["Exception Report"]
    assert ws["A1"].value == "Category" and ws["A2"].value == "No exceptions"
